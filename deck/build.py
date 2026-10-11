"""Fill the official AI Builder Cup template → deck/VibeRick-Dub.pptx (+ PDF via LibreOffice).

uv run --with python-pptx python deck/build.py
Baseline run metrics come from runs/. Other historical comparisons remain explicitly labeled baseline.
Use --output or DUB_DECK_OUTPUT to preserve the original deck when making a review draft."""
import argparse
import os
import copy
import json
import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = ROOT.parent / "submission-template.pptx"
A = HERE / "assets"
OUT = Path(os.environ.get("DUB_DECK_OUTPUT", HERE / "VibeRick-Dub.pptx"))
RUN = os.environ.get("DUB_DECK_RUN", "Fty2sz3dBwU.180s")

INK, MUTED, ACCENT = RGBColor(0x11, 0x12, 0x14), RGBColor(0x5A, 0x5D, 0x66), RGBColor(0x2F, 0x5B, 0xFF)
FONT = "Arial"
L, TOP, W, BOTTOM = Emu(265350), Emu(1180000), Emu(8613300), Emu(5000000)

LINKS = {
    "product": "https://viberick-dub.web.app",
    "github": "Public repository: pending publication",
    "video": "Demo video URL: pending approval and upload",
}


def report(name: str) -> dict | None:
    p = ROOT / "runs" / name / "out" / "report.json"
    return json.loads(p.read_text()) if p.exists() else None


def mmss(s: float) -> str:
    return f"{int(s // 60)}:{int(s % 60):02d}"


def set_title(slide, text: str) -> None:
    box = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    tf = box.text_frame
    p0 = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    run = p0.runs[0] if p0.runs else p0.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(20), True, FONT
    run.font.color.rgb = INK
    box.height = Emu(480000)


def text(slide, x, y, w, h, paras: list, size: int = 13, gap: int = 6):
    """paras: str | (str, dict(size, bold, color, bullet))"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(paras):
        s, o = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(o.get("gap", gap))
        r = p.add_run()
        r.text = ("•  " if o.get("bullet") else "") + s
        r.font.name, r.font.size = FONT, Pt(o.get("size", size))
        r.font.bold = o.get("bold", False)
        r.font.color.rgb = o.get("color", INK)
    return tb


def image(slide, path: Path, x, y, w=None, h=None):
    from PIL import Image
    iw, ih = Image.open(path).size
    if w and h:
        if iw / ih > w / h:
            h2 = int(w * ih / iw)
            y, h = y + (h - h2) // 2, h2
        else:
            w2 = int(h * iw / ih)
            x, w = x + (w - w2) // 2, w2
    return slide.shapes.add_picture(str(path), x, y, w, h)


def table(slide, x, y, w, rows: list[list[str]], col_w: list[float], size: int = 11, row_h: int = 300000):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Emu(row_h * len(rows))).table
    for j, f in enumerate(col_w):
        t.columns[j].width = Emu(int(w * f))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = val
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0x11, 0x12, 0x14) if i == 0 else (
                RGBColor(0xF3, 0xF6, 0xFF) if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name, r.font.size = FONT, Pt(size)
                    r.font.bold = i == 0
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
            c.margin_left = c.margin_right = Emu(70000)
            c.margin_top = c.margin_bottom = Emu(30000)
    return t


def drop(prs, index: int) -> None:
    sld = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)


def build() -> Path:
    demo, full, multi = report(RUN), report("Fty2sz3dBwU"), report("PqVbypvxDto.120s")
    if demo is None:
        raise FileNotFoundError(f"Missing baseline report for run {RUN!r}; refusing placeholder metrics")
    prs = Presentation(str(TEMPLATE))
    s = prs.slides

    # 1 · Team details
    box = next(sh for sh in s[0].shapes if sh.has_text_frame)
    tf = box.text_frame
    lines = ["VibeRick Dub", "", "Team name: VibeRick", "Team leader name: Jaeho Shin",
             "Problem Statement: Media, Content & Digital Experiences — creators reach only the viewers who "
             "speak their language. VibeRick Dub automates multilingual audio, subtitle timing and output delivery."]
    for i, p in enumerate(tf.paragraphs):
        for r in p.runs:
            r.text = ""
        if i < len(lines) and p.runs:
            p.runs[0].text = lines[i]
    while len(tf.paragraphs) < len(lines):
        p = tf.add_paragraph()
        p.add_run().text = lines[len(tf.paragraphs) - 1]
    image(s[0], A / "logo-v.png", Emu(7900000), Emu(3150000), Emu(950000), Emu(950000))

    # 3 · Outcome first; detailed proof follows.
    set_title(s[2], "You made the video. We handle the languages.")
    text(s[2], L, TOP, Emu(5250000), Emu(3750000), [
        ("One recording. Multilingual versions.", {"size": 22, "bold": True, "gap": 12}),
        "Paste a video link and choose languages. VibeRick Dub translates for the available speaking time, generates audio, checks each take and exports video, audio and synced subtitles.",
        "Creators can use a consented voice replica. Other speakers use distinct presets. Language switching keeps the same playback position; line editing is available when needed.",
        ("Four tested pipeline languages: Korean, English, Japanese and Simplified Chinese. Automatic checks are not human quality certification.", {"color": MUTED, "size": 11}),
    ], size=14, gap=10)
    text(s[2], Emu(6000000), Emu(TOP + 300000), Emu(2800000), Emu(3100000), [
        ("Finished outputs", {"size": 16, "color": MUTED, "gap": 20}),
        ("Dubbed video", {"size": 24, "bold": True, "gap": 20}),
        ("Audio tracks", {"size": 24, "bold": True, "gap": 20}),
        ("Synced subtitles", {"size": 24, "bold": True, "gap": 20}),
        ("Play and download each language. Fine-tune only when needed.", {"size": 13, "color": MUTED}),
    ])

    # 4 · Problem, differentiation and commercial hypothesis.
    set_title(s[3], "A finished language version, without coordinating every step")
    text(s[3], L, TOP, W, Emu(3800000), [
        ("The work repeats for every language", {"bold": True, "size": 17}),
        "Translation, recording, timing, subtitle preparation and quality checks form a connected production task. Our first users are independent tutorial creators, podcasters and small course teams.",
        ("The system measures and repairs its own output", {"bold": True, "size": 17}),
        "Translation length is a starting estimate. Actual speech duration and model listening checks drive retries; unresolved issues remain visible. Outputs can be downloaded and edited before creator-controlled publishing.",
        ("A focused route to a business", {"bold": True, "size": 17}),
        "Test paid usage per output minute with creators who localize repeatedly. Measure active labor, generation cost and the share of lines needing intervention before setting pricing.",
        ("YouTube documents that automatic dubs cannot be edited. We do not claim superior audio quality without a matched comparison. Source: support.google.com/youtube/answer/15569972", {"size": 10, "color": MUTED}),
    ], size=14, gap=10)

    # 5 · Features
    set_title(s[4], "List of features offered by the solution")
    feats = [
        ("Distinct speaker voices", "Current demo: Charon for Lex, Algieba for Demis. Consented voice replication is a separate tested feature."),
        ("Translation that fits the clock", "Each sentence gets a character budget from its time slot and the voice's measured speed"),
        ("Automatic audio checks and repair", "Checks language, content and source meaning; retries or flags unresolved issues."),
        ("Audio-led transcription", "Chirp timestamps + Gemini audio proofreading; source captions can contain editorial edits."),
        ("Multiple speakers", "Speaker labels and separate voice assignments; the podcast excerpt uses transcript-assisted source alignment."),
        ("Consent lock", "Clone only for the channel owner whose voice verifies against the consent recording"),
        ("Play, download, optionally edit", "Same-position language switching, per-language outputs and preview / accept / discard"),
        ("Ready to upload", "Dubbed MP4 + M4A (−14 LUFS, original music kept) + SRT/VTT subtitles"),
    ]
    for k, (h, d) in enumerate(feats):
        col, row = k % 2, k // 2
        text(s[4], Emu(265350 + col * 4350000), Emu(1180000 + row * 930000), Emu(4200000), Emu(880000),
             [(h, {"bold": True, "size": 14, "gap": 3}), (d, {"color": MUTED, "size": 12})])

    # 6 · Process flow
    set_title(s[5], "Process flow diagram")
    image(s[5], A / "flow.png", L, TOP, W, Emu(3800000))

    # 7 · Actual outcome-first UI, preserved aspect ratio.
    set_title(s[6], "The finished video leads the experience")
    image(s[6], ROOT / "docs/review-2026-10-10/outcome-studio.jpg", L, TOP, Emu(4600000), Emu(3650000))
    text(s[6], Emu(5100000), TOP, Emu(3600000), Emu(3650000), [
        ("Listen. Switch languages. Download.", {"bold": True, "size": 20}),
        "A large player, aligned captions and language-specific video, audio and subtitle files make the result immediately usable.",
        "Detailed timing, transcript and quality records are expandable. Optional edits preserve the user's exact words, with preview before acceptance.",
        ("Actual local UI capture. Outcome-first layout is awaiting deployment; the previous studio is deployed. This screenshot uses a recorded creator sample.", {"size": 11, "color": MUTED}),
    ], size=14, gap=12)

    # 8 · Architecture
    set_title(s[7], "Architecture diagram of the proposed solution")
    image(s[7], A / "arch.png", L, TOP, W, Emu(3800000))

    # 9 · Technologies
    set_title(s[8], "Technologies used in the solution")
    table(s[8], L, TOP, W, [
        ["Technology", "Role"],
        ["Gemini 3.8 Flash (Vertex AI)", "Proofread by listening · dubbing brief · translate to time · shorten · judge every clip · cast voices"],
        ["Gemini 3.8 Flash TTS — voice replication", "The creator's own voice in other languages (SynthID watermarked)"],
        ["Chirp 3 (Speech-to-Text v2)", "Word timestamps + speaker diarization — the timing backbone"],
        ["Cloud Run (service + Job)", "FastAPI API · dubbing worker (4 vCPU, demucs, ECAPA, ffmpeg)"],
        ["Firestore", "Jobs, live agent event stream, per-line results (realtime listeners)"],
        ["Firebase Hosting", "Web app; /api/** rewritten to Cloud Run"],
        ["Cloud Storage", "Batch STT input; public dubbed MP4/M4A/SRT/VTT"],
        ["Secret Manager · Artifact Registry · Cloud Build", "Keys · container image · builds"],
        ["Open source", "demucs (music bed), ECAPA-TDNN (speaker verification for the consent lock), yt-dlp, ffmpeg"],
    ], [0.36, 0.64], size=11, row_h=330000)

    # 10 · Cost: only the text usage comes from the recorded report.
    set_title(s[9], "Estimated implementation cost")
    usd = demo.get("gemini_text_usd") if demo else None
    dubbed_min = (demo["duration"] / 60 * len(demo["langs"])) if demo else 1
    text_pm = usd / dubbed_min if usd is not None else None
    table(s[9], L, TOP, W, [
        ["Selected run: cost per dubbed minute", "USD", "Evidence"],
        ["Gemini text + listening checks", f"{text_pm:.3f}" if text_pm is not None else "Pending", "Usage recorded; rate-based estimate" if text_pm is not None else "Aggregation pending"],
        ["Gemini TTS", "0.014", "Planning estimate"],
        ["Chirp (source minutes split across 3 languages)", "0.005", "Planning estimate"],
        ["Cloud Run compute", "0.003", "Planning estimate"],
        ["Illustrative subtotal", f"~ {text_pm + 0.022:.2f}" if text_pm is not None else "Pending", "Not a billing reconciliation"],
    ], [0.57, 0.13, 0.30], size=12, row_h=360000)
    text(s[9], L, Emu(3600000), W, Emu(1150000), [
        "Storage, downloads/proxy, retries and additional line previews can add cost. "
        "Fresh aggregate usage and actual billing are still being reconciled; no revised cost claim is made.",
        ("Infrastructure uses Google Cloud. Voice generation uses a separately billed Google project; "
         "credits and cash billing must be tracked separately.", {"color": MUTED, "size": 11})], size=13)

    # 11 · Current podcast run evidence.
    set_title(s[10], "A real two-speaker dubbing run")
    podcast = report("podcast-demo-2026-10-11")
    passed = sum(row["ok"] for row in podcast["langs"])
    total = sum(row["lines"] for row in podcast["langs"])
    table(s[10], L, TOP, W, [
        ["Recorded demo", "Evidence"],
        ["Source", "Lex Fridman Podcast #475 · Demis Hassabis · 56:27–57:07"],
        ["Speaker assignments", "Lex: Charon preset · Demis: Algieba preset"],
        ["Target tracks", "Korean · Japanese · Chinese (Simplified)"],
        ["Automatic listening checks", f"{passed}/{total} passed; this is not native listening certification"],
        ["Source timing", "Audio-estimated alignment against the published transcript"],
        ["Later editorial repair", "Korean opening terminology revised and rechecked"],
    ], [0.32, 0.68], size=12, row_h=390000)
    text(s[10], L, Emu(4100000), W, Emu(700000), [
        ("The recorded run and later editorial repair are preserved separately. No end-to-end Cloud Run speed or "
         "automatic diarization claim is inferred from this selected demo.", {"size": 11, "color": MUTED})])

    # 12 · Fresh comparative evidence, measured over the same 180 seconds.
    set_title(s[11], "Separate creator test: measured pauses over 3 minutes")
    before_doc = json.loads((ROOT / "docs/review-2026-10-10/quality-fresh-3min-2026-10-11.json").read_text())
    after_doc = json.loads((ROOT / "runs/refined-3min-2026-10-11/refinement-quality.json").read_text())
    before = next(a for a in before_doc["audits"] if Path(a["run"]).name == "Fty2sz3dBwU.180s")
    after = next(a for a in after_doc["audits"] if Path(a["run"]).name == "refined-3min-2026-10-11")
    def longest(audio):
        return max((g["seconds"] for g in audio["gaps"] if g["source_active_seconds"] > 0), default=0)
    rows = [["Language", "Detected silence: old → refined", "Longest gap*: old → refined"]]
    for lang, label in [("en", "English"), ("ja", "Japanese"), ("zh-Hans", "Chinese")]:
        old, new = before["languages"][lang]["audio"], after["languages"][lang]["audio"]
        assert old["comparison_window_seconds"] == new["comparison_window_seconds"] == 180
        rows.append([label, f"{old['dub_detected_silence_seconds']:.2f} → {new['dub_detected_silence_seconds']:.2f} s",
                     f"{longest(old):.2f} → {longest(new):.2f} s"])
    table(s[11], L, TOP, W, rows, [0.19, 0.40, 0.41], size=12, row_h=390000)
    refined = report("refined-3min-2026-10-11")
    ok = sum(l["ok"] for l in refined["langs"])
    n = sum(l["lines"] for l in refined["langs"])
    trials = json.loads((ROOT / "runs/refined-3min-2026-10-11/refinement-results.json").read_text())
    accepted = sum(t["accepted"] is True for t in trials)
    text(s[11], L, Emu(2900000), W, Emu(1900000), [
        ("Separate Korean creator recording: RMS energy at −40 dBFS over the same 180 s. This is an energy heuristic, not a test of naturalness or intelligibility.", {"size": 12}),
        ("*Dub-silent intervals with any source-active overlap. Music or noise can affect detection.", {"size": 11, "color": MUTED}),
        (f"Refined run: {ok}/{n} automatic checks passed; {accepted}/{len(trials)} candidate line refinements retained. Native listening review remains pending.", {"size": 12}),
        ("Chinese longest gap increased 2.82 → 2.94 s vs the first revision. Improvements are not universal.", {"size": 11, "color": MUTED}),
        ("Separate historical 15:05 baseline: 366/369 checks passed; 3 lines flagged. Fresh generation ran as serial local validation, not a Cloud Run latency benchmark.", {"size": 11, "color": MUTED}),
    ], gap=7)

    # 13 · Architecture depth and evidence still needed.
    set_title(s[12], "Operate reliably, then expand language coverage")
    text(s[12], L, TOP, W, Emu(3800000), [
        ("Implemented reliability controls", {"bold": True, "size": 16}),
        "Shared preprocessing, bounded language concurrency and independent language completion. Idempotent job creation, owner-protected edits, durable workspace snapshots, worker leases and usage caps.",
        ("Verified cloud behavior", {"bold": True, "size": 16}),
        "Recorded smoke tests cover job creation, duplicate-request recovery, unauthorized edit rejection, preview acceptance/discard and exact-wording preservation. API revision 00010-rih was promoted after validation.",
        ("Scale by measuring the bottlenecks", {"bold": True, "size": 16}),
        "Next: sustained load tests, quota/backpressure tuning, per-stage cost accounting and retention enforcement. Wider language support requires separate listening and timing validation.",
        ("Release gates: representative media rights, human listening, matched timing/cost evidence and final public repository/video links. New UI is local; no scale or hours-saved claim is made.", {"color": MUTED, "size": 11}),
    ], size=14, gap=10)

    # 14 · Links
    set_title(s[13], "Links")
    text(s[13], L, TOP, W, Emu(2400000), [
        ("GitHub public repository", {"bold": True}), (LINKS["github"], {"color": ACCENT}),
        ("Demo video (under 3 minutes)", {"bold": True}), (LINKS["video"], {"color": ACCENT}),
        ("Final product", {"bold": True}), (LINKS["product"], {"color": ACCENT}),
    ], size=15, gap=6)

    # 15 · Responsible AI (blank slide in the template)
    tb = text(s[14], L, Emu(663825), W, Emu(480000), [("Consent and safety by design", {"bold": True, "size": 20})])
    text(s[14], L, TOP, Emu(5000000), Emu(3800000), [
        ("Two locks before a voice is cloned", {"bold": True, "size": 14}),
        ("1. Channel lock: the channel owner recorded the consent statement.", {"size": 13}),
        ("2. Speaker lock: the diarized speaker's voice verifies against that recording (ECAPA cosine ≥ 0.45).",
         {"size": 13}),
        ("Everyone else gets a Gemini preset voice chosen by Gemini after listening to them.", {"size": 13}),
        ("Verification stores an embedding. Voice replication also uses provider-side voice state; local consent/reference "
         "recordings and generated media need explicit retention controls.",
         {"size": 13, "color": MUTED}),
    ], gap=9)
    table(s[14], Emu(5500000), TOP, Emu(3380000), [
        ["Boundary", "Behavior"],
        ["Consent absent", "Use preset"],
        ["Speaker match ambiguous", "Use preset"],
        ["Automated checks fail", "Retry or flag"],
        ["Manual edit", "Preserve wording"],
        ["Public output", "Disclosed before creation"],
    ], [0.50, 0.50], size=11, row_h=430000)

    sources = {
        3: "YouTube automatic dubbing documentation checked 2026-10-11: https://support.google.com/youtube/answer/15569972 . User/labor benefit is a hypothesis pending measured workflow comparison.",
        2: "The main screenshot is an actual local creator result. Separate podcast evidence uses official Lex Fridman Podcast #475, https://www.youtube.com/watch?v=-HzgcbRXUK8 , selected 56:27–57:07. Actual local frames from runs/podcast-demo-2026-10-11/source.mp4. Charon and Algieba are preset voices; these are not cloned celebrity voices. Source-alignment.json documents transcript-assisted alignment.",
        6: "Actual local UI capture docs/review-2026-10-10/outcome-studio.jpg. The current outcome-first layout is local, not proof of a deployed revision.",
        10: "runs/podcast-demo-2026-10-11/out/report.json; units.json; source-alignment.json; revisions/ko0-before. All18 target rows passed automatic checks. Local run timing is not advertised as deployment latency.",
        9: f"Baseline: runs/{RUN}/out/report.json. gemini_text_usd is token usage priced at configured rates in dub/gem.py, not an invoice. Other rows are planning estimates from the previous deck.",
        11: "Measured comparison: docs/review-2026-10-10/quality-fresh-3min-2026-10-11.json and runs/refined-3min-2026-10-11/refinement-quality.json, RMS threshold -40 dBFS, 180-second common window. Candidate acceptance: runs/refined-3min-2026-10-11/refinement-results.json. Chinese longest gap regressed from 2.82 to 2.94 seconds relative to the first revision, despite improvement against original 6.56 seconds. Energy is not native listening quality. Historical full-run baseline retains 3 flagged lines. Fresh sec_total is local serial-validation artifact timing, not deployed Cloud Run latency.",
        14: "Implementation: dub/speakers.py, dub/config.py. Reference/consent audio and provider voice state require separate retention treatment.",
    }
    for idx, source in sources.items():
        s[idx].notes_slide.notes_text_frame.text = source
    drop(prs, 1)   # the template's "how to download" instruction slide
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)],
                   check=True, capture_output=True)
    return OUT


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run", default=RUN, help="Existing runs/ directory used for baseline metrics")
    parser.add_argument("--output", type=Path, default=OUT, help="Separate review PPTX path (PDF written beside it)")
    args = parser.parse_args()
    RUN, OUT = args.run, args.output.resolve()
    print(build())
