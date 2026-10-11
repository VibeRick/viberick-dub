"""Deck v3 on the official template → deck/v3/VibeRick-Dub.pptx + PDF.

uv run --with python-pptx --with pillow python deck/build_v3.py
Keeps every template heading; visuals from deck/v3/assets (UI captures of the deployed studio, diagrams).
Business figures: sources in speaker notes (checked 2026-10-11)."""
import json
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import build as B  # noqa: E402  (template helpers)

ROOT = HERE.parent
A = HERE / "v3" / "assets"
OUT = HERE / "v3" / "VibeRick-Dub.pptx"
INK, MUTED, ACC = RGBColor(0x11, 0x12, 0x14), RGBColor(0x5F, 0x5C, 0x57), RGBColor(0xE8, 0x55, 0x1A)
L, TOP, W = B.L, B.TOP, B.W
LINKS = json.loads((HERE / "v3" / "links.json").read_text()) if (HERE / "v3" / "links.json").exists() else {
    "github": "github.com/… (publishing)", "video": "youtu.be/… (uploading)", "product": "https://viberick-dub.web.app"}


def t(slide, x, y, w, h, paras, size=13, gap=6):
    return B.text(slide, x, y, w, h, paras, size=size, gap=gap)


def stat(slide, x, y, w, big, small, color=INK):
    t(slide, x, y, w, Emu(900000), [(big, {"size": 30, "bold": True, "color": color, "gap": 2}), (small, {"size": 11, "color": MUTED})])


def build():
    full, demo = B.report("Fty2sz3dBwU"), B.report("Fty2sz3dBwU.180s")
    prs = Presentation(str(B.TEMPLATE))
    s = prs.slides

    # 0 · Team
    box = next(sh for sh in s[0].shapes if sh.has_text_frame)
    tf = box.text_frame
    lines = ["VibeRick Dub", "", "Team name: VibeRick", "Team leader name: Jaeho Shin",
             "Problem Statement: Media, Content & Digital Experiences — creators only reach viewers who speak "
             "their language. We dub their videos in their own voice."]
    for i, p in enumerate(tf.paragraphs):
        for r in p.runs:
            r.text = ""
        if i < len(lines) and p.runs:
            p.runs[0].text = lines[i]
    while len(tf.paragraphs) < len(lines):
        tf.add_paragraph().add_run().text = lines[len(tf.paragraphs) - 1]
    B.image(s[0], B.A / "logo-v.png", Emu(7900000), Emu(3150000), Emu(950000), Emu(950000))

    # 2 · Brief about the idea
    t(s[2], L, TOP, Emu(3900000), Emu(3700000), [
        ("Your voice. Every language.", {"size": 26, "bold": True, "gap": 12}),
        "Paste a YouTube link and pick from 125 languages. An agent on Google Cloud transcribes it, translates each sentence to fit "
        "the moment it is spoken, says it in the creator's own cloned voice, then listens to every line and repairs what is wrong.",
        "Out come dubbed MP4s, M4A audio tracks ready for YouTube's multi-language audio, and SRT/VTT subtitles. "
        "Creators review any line, edit it, and regenerate only that line.",
        ("Live: viberick-dub.web.app", {"bold": True, "color": ACC}),
    ], size=13, gap=10)
    B.image(s[2], A / "ui-hero.png", Emu(4350000), TOP, Emu(4530000), Emu(3500000))

    # 3 · Opportunities / differentiation / USP
    B.set_title(s[3], "Opportunities · how it is different · USP")
    t(s[3], L, TOP, Emu(4200000), Emu(3800000), [
        ("The opportunity", {"bold": True, "size": 16, "color": ACC}),
        "Creators who add multi-language audio get over 25% of watch time from views in the video's non-primary language (YouTube, Sep 2025).",
        "More than 6 million people a day already watch 10+ minutes of auto-dubbed video (YouTube, Dec 2025).",
        ("What's missing", {"bold": True, "size": 16, "color": ACC}),
        "YouTube can auto-dub a Korean video only into English or Indonesian, the dubs cannot be edited, and they are not the creator's voice.",
        "Dubbing tools charge about $0.50–0.67 per dubbed minute (ElevenLabs, Perso, Rask list prices, Oct 2026).",
    ], size=12, gap=7)
    t(s[3], Emu(4650000), TOP, Emu(4230000), Emu(3800000), [
        ("How VibeRick Dub is different", {"bold": True, "size": 16, "color": ACC}),
        ("125 target languages: every language Gemini 3.8 TTS speaks, Hindi, Tamil and Bangla included.", {"bullet": True}),
        ("Own voice, consent-gated: two locks (channel consent + speaker verification) before a clone is used.", {"bullet": True}),
        ("An agent that listens: Gemini hears every generated line and repairs wrong language, missing words and overruns.", {"bullet": True}),
        ("Timed to the speaker: translation is budgeted to each sentence's slot and the voice's measured speaking rate.", {"bullet": True}),
        ("Inspectable and editable: every agent decision is visible; edit one line, preview, accept.", {"bullet": True}),
        ("≈ $0.03 per dubbed minute to run, so a price far below today's tools still leaves a strong margin.", {"bullet": True}),
    ], size=12, gap=7)

    # 4 · Features
    feats = [
        ("Creator's own voice", "Gemini TTS voice replication from a 20 s reference, used only behind two consent locks."),
        ("Every speaker, own voice", "Chirp 3 tells speakers apart; each gets the creator's clone (with consent) or the closest preset Gemini picks by ear."),
        ("Translation that fits the clock", "Per-sentence character budgets from the slot and the voice's measured speed."),
        ("Listen-back repair loop", "Measure → shorten → time-fit → Gemini listens → re-speak → re-translate → flag."),
        ("Live studio", "Agent events stream from Firestore; switch languages at the same playback position."),
        ("Line edit with preview", "Edit one line, hear the preview in the same voice, accept or discard."),
        ("Upload-ready outputs", "MP4 + M4A (−14 LUFS over the original music bed) + SRT/VTT per language."),
        ("125 languages", "Any target Gemini 3.8 TTS speaks, up to 4 per video; Korean, English, Japanese, Chinese sources."),
    ]
    for k, (h, d) in enumerate(feats):
        col, row = k % 2, k // 2
        t(s[4], Emu(265350 + col * 4350000), Emu(1180000 + row * 900000), Emu(4200000), Emu(860000),
          [(h, {"bold": True, "size": 14, "gap": 3}), (d, {"color": MUTED, "size": 12})])

    # 5 · Process flow
    B.image(s[5], A / "flow.png", L, TOP, W, Emu(2400000))
    t(s[5], L, Emu(3700000), W, Emu(900000), [
        ("Every step writes an event to Firestore, so the creator watches the agent decide in real time. "
         "Unresolved lines stay flagged in the studio instead of being silently shipped.", {"size": 13, "color": MUTED})])

    # 6 · Wireframes → the real studio
    B.image(s[6], A / "ui-studio.png", L, TOP, Emu(5300000), Emu(3600000))
    t(s[6], Emu(5750000), TOP, Emu(3130000), Emu(3600000), [
        ("The studio", {"bold": True, "size": 18}),
        "Player with same-position language switching, synced transcript, live agent activity, a timeline of every line in every language, and a line inspector.",
        ("Deployed UI, not a mock: viberick-dub.web.app", {"size": 12, "color": ACC}),
    ], size=13, gap=10)

    # 7 · Architecture
    B.image(s[7], A / "arch.png", L, TOP, W, Emu(3700000))

    # 8 · Technologies
    B.table(s[8], L, TOP, W, [
        ["Technology", "Role in VibeRick Dub"],
        ["Gemini 3.8 Flash (Vertex AI)", "Proofreads by listening · dubbing brief · translate to time · shorten · listens to every clip · casts voices"],
        ["Gemini TTS voice replication", "The creator's own voice in other languages; presets for other speakers (SynthID watermarked)"],
        ["Chirp 3 (Speech-to-Text v2)", "Word timestamps + speaker diarization: the timing backbone"],
        ["Cloud Run service + Cloud Run Jobs", "FastAPI API; one dubbing worker execution per job (4 vCPU)"],
        ["Firestore", "Jobs, live agent event stream, per-line results, edits (real-time listeners)"],
        ["Firebase Hosting", "The studio; /api/** rewritten to Cloud Run"],
        ["Cloud Storage", "STT input; public MP4/M4A/SRT/VTT outputs with lifecycle expiry"],
        ["Secret Manager · Artifact Registry · Cloud Build", "Keys · container image · builds"],
        ["Open source", "demucs (music bed) · ECAPA-TDNN (consent voice match) · ffmpeg · yt-dlp"],
    ], [0.36, 0.64], size=11, row_h=330000)

    # 9 · Cost + business model
    usd = demo["gemini_text_usd"] / (demo["duration"] / 60 * len(demo["langs"]))
    B.table(s[9], L, TOP, Emu(4300000), [
        ["Cost per dubbed minute", "USD"],
        ["Gemini text + listening checks (recorded usage)", f"{usd:.3f}"],
        ["Gemini TTS (estimate)", "0.014"],
        ["Chirp 3 (estimate)", "0.005"],
        ["Cloud Run compute (estimate)", "0.003"],
        ["Total", f"≈ {usd + 0.022:.2f}"],
    ], [0.78, 0.22], size=12, row_h=380000)
    t(s[9], Emu(4800000), TOP, Emu(4080000), Emu(3800000), [
        ("Business model", {"bold": True, "size": 16, "color": ACC}),
        ("Pay per dubbed minute, plus a creator subscription with monthly minutes.", {"bullet": True}),
        ("At $0.20 per dubbed minute: less than half of today's $0.50–0.67 tools, ≈ 85% gross margin.", {"bullet": True}),
        ("Infra scales to zero: each dub is one Cloud Run Job execution; an idle service costs nothing.", {"bullet": True}),
        ("First market: Korean creators. YouTube says about 30% of Korean channels' watch time already comes from overseas (Aug 2026).", {"bullet": True}),
        ("Estimates exclude storage egress and proxy download costs.", {"size": 10, "color": MUTED}),
    ], size=12, gap=7)

    # 10 · Snapshots
    B.image(s[10], A / "ui-lines.png", L, TOP, W, Emu(3300000))
    t(s[10], L, Emu(4500000), W, Emu(400000), [
        ("Line review and inspector, showing the real repair history of one line: “Said ‘WhyBrick’ instead of ‘VibeRick’” → generated again → passed. "
         "The timeline and live agent log are on the same page.", {"size": 11, "color": MUTED})])

    # 11 · Performance / benchmarking
    ok3, n3 = sum(l["ok"] for l in demo["langs"]), sum(l["lines"] for l in demo["langs"])
    okf, nf = sum(l["ok"] for l in full["langs"]), sum(l["lines"] for l in full["langs"])
    stat(s[11], L, TOP, Emu(2100000), f"{okf}/{nf}", "lines passed Gemini's listening check · full 15:05 video, 3 languages", ACC)
    stat(s[11], Emu(2450000), TOP, Emu(2100000), f"{ok3}/{n3}", "3 min × 3 languages, end to end in 2:52")
    stat(s[11], Emu(4600000), TOP, Emu(2100000), "≤ 0.16 s", "p95 line end overrun vs the original speech")
    stat(s[11], Emu(6750000), TOP, Emu(2100000), "0.90", "creator voice match · others ≤ 0.09 · gate 0.45")
    B.table(s[11], L, Emu(2350000), W, [
        ["Transcription (15:05 Korean video, 3 runs)", "Character error", "Timing |dt| p50"],
        ["Gemini 3.8 Flash alone", "10.6–11.8%", "5–12 s (drifts up to 42 s)"],
        ["Chirp 3 alone", "8.6%", "0.22 s"],
        ["Chirp 3 + Gemini proofread (ours)", "7.6–7.9%", "0.22–0.24 s"],
    ], [0.5, 0.22, 0.28], size=12, row_h=360000)
    t(s[11], L, Emu(3900000), W, Emu(900000), [
        ("Deployed run, 2026-10-11: a two-speaker Google DeepMind podcast excerpt (3 min, English) into Korean, Japanese and Chinese "
         "on Cloud Run Jobs: 115/117 lines passed after repairs, 2 flagged for a human, 9 min 8 s from submit to finished MP4s.", {"size": 12, "gap": 4}),
        ("STT reference: the creator's own edited captions. Automated listening checks are not a native-speaker rating.", {"size": 10, "color": MUTED})])

    # 12 · Future
    t(s[12], L, TOP, Emu(4200000), Emu(3800000), [
        ("Next 3 months", {"bold": True, "size": 16, "color": ACC}),
        ("One-click publish of audio tracks to YouTube's multi-language audio.", {"bullet": True}),
        ("More languages, each added only after timing and listening validation.", {"bullet": True}),
        ("Glossary per channel: names and terms the creator fixes once, applied to every video.", {"bullet": True}),
        ("Longer videos: chunked Chirp batches and per-language parallel workers.", {"bullet": True}),
    ], size=13, gap=8)
    t(s[12], Emu(4650000), TOP, Emu(4230000), Emu(3800000), [
        ("Scale & operations", {"bold": True, "size": 16, "color": ACC}),
        ("Each dub is an independent Cloud Run Job execution: horizontal by design.", {"bullet": True}),
        ("Quotas, idempotent creation and owner tokens already guard the public API.", {"bullet": True}),
        ("Per-stage cost ledger → pricing per minute with a known margin.", {"bullet": True}),
        ("Consent registry for many creators: one verified voice per enrolled speaker.", {"bullet": True}),
    ], size=13, gap=8)

    # 13 · Links
    B.set_title(s[13], "Provide links to your: GitHub Public Repository · Demo Video · Final Product")
    t(s[13], L, Emu(1700000), W, Emu(2600000), [
        ("GitHub public repository", {"bold": True}), (LINKS["github"], {"color": ACC, "gap": 14}),
        ("Demo video (under 3 minutes)", {"bold": True}), (LINKS["video"], {"color": ACC, "gap": 14}),
        ("Final product", {"bold": True}), (LINKS["product"], {"color": ACC}),
    ], size=15, gap=4)

    # 14 · Consent and safety (blank template slide)
    t(s[14], L, Emu(663825), W, Emu(480000), [("Consent and safety by design", {"bold": True, "size": 20})])
    t(s[14], L, TOP, Emu(4400000), Emu(3800000), [
        ("Two locks before a voice is cloned", {"bold": True, "size": 15, "color": ACC}),
        "1. Channel lock: the video's channel belongs to a creator who recorded a spoken consent statement.",
        "2. Speaker lock: the diarized speaker verifies against that recording (ECAPA cosine ≥ 0.45; creator 0.90, others ≤ 0.09).",
        "Everyone else gets a preset voice that Gemini picks after listening to them. Dubbed audio carries Gemini's SynthID watermark.",
    ], size=13, gap=9)
    B.table(s[14], Emu(4900000), TOP, Emu(3980000), [
        ["Situation", "What the agent does"],
        ["No consent on file", "Preset voice"],
        ["Speaker match ambiguous", "Preset voice"],
        ["Listening check fails", "Retry, then flag"],
        ["Creator edits a line", "Exact wording kept; preview first"],
        ["Public outputs", "Expire after 7 days"],
    ], [0.48, 0.52], size=12, row_h=420000)

    notes = {
        3: "YouTube multi-language audio: blog.youtube/news-and-events/multi-language-audio/ (2025-09-10). Auto-dub viewers and 27 languages: blog.google/intl/en-in/products/platforms/unlocking-a-global-audience-with-auto-dubbing/ (2026-02-05). Korean → English/Indonesian only; cannot edit: support.google.com/youtube/answer/15569972. Prices: elevenlabs.io/pricing (Pro $99/600K credits, 3,000 credits/min), perso.ai/en/pricing (Pro $99 / 180 min), rask.ai/pricing (Creator Pro annual $799 / 1,200 min); per-minute values computed 2026-10-11.",
        9: "Gemini text cost is recorded token usage priced at configured rates (dub/gem.py), not an invoice. Korean overseas watch share: view.asiae.co.kr/en/article/2026081014010865557 (YouTube VP, paraphrased, 2026-08-10). Market estimate if asked: Valuates, AI dubbing USD 783M (2023) → 1.88B (2030), via Speech Technology Magazine.",
        11: "runs/Fty2sz3dBwU/out/report.json (full), runs/Fty2sz3dBwU.180s/out/report.json (3 min). STT: RESEARCH-2026-10-10.md, experiments/2026-10-10-stt/score.txt. ECAPA: voice/ecapa.py.",
    }
    for i, n in notes.items():
        s[i].notes_slide.notes_text_frame.text = n
    B.drop(prs, 1)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT.parent), str(OUT)], check=True, capture_output=True)
    return OUT


if __name__ == "__main__":
    print(build())
