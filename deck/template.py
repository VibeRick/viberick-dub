"""Helpers for filling the official AI Builder Cup slide template (python-pptx)."""
import json
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
TEMPLATE = ROOT.parent / "submission-template.pptx"  # the organiser's template; not part of this repo
A = HERE / "v3" / "assets"

INK = RGBColor(0x11, 0x12, 0x14)
FONT = "Arial"
L, TOP, W, BOTTOM = Emu(265350), Emu(1180000), Emu(8613300), Emu(5000000)


def report(name: str) -> dict | None:
    p = ROOT / "runs" / name / "out" / "report.json"
    return json.loads(p.read_text()) if p.exists() else None


def mmss(s: float) -> str:
    return f"{int(s // 60)}:{int(s % 60):02d}"


def set_title(slide, text: str) -> None:
    box = next(sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip())
    tf = box.text_frame
    p0 = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs)[1:]:
        r._r.getparent().remove(r._r)
    run = p0.runs[0] if p0.runs else p0.add_run()
    run.text = text
    run.font.size, run.font.bold, run.font.name = Pt(20), True, FONT
    run.font.color.rgb = INK
    box.height = Emu(480000)


def text(slide, x, y, w, h, paras: list, size: int = 13, gap: int = 6):
    """paras: str | (str, dict(size, bold, color, bullet))"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(paras):
        s, o = (item, {}) if isinstance(item, str) else item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(o.get("gap", gap))
        r = p.add_run()
        r.text = ("•  " if o.get("bullet") else "") + s
        r.font.name, r.font.size = FONT, Pt(o.get("size", size))
        r.font.bold = o.get("bold", False)
        r.font.color.rgb = o.get("color", INK)
    return tb


def image(slide, path: Path, x, y, w=None, h=None):
    from PIL import Image
    iw, ih = Image.open(path).size
    if w and h:
        if iw / ih > w / h:
            h2 = int(w * ih / iw)
            y, h = y + (h - h2) // 2, h2
        else:
            w2 = int(h * iw / ih)
            x, w = x + (w - w2) // 2, w2
    return slide.shapes.add_picture(str(path), x, y, w, h)


def table(slide, x, y, w, rows: list[list[str]], col_w: list[float], size: int = 11, row_h: int = 300000):
    t = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Emu(row_h * len(rows))).table
    for j, f in enumerate(col_w):
        t.columns[j].width = Emu(int(w * f))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.text = val
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0x11, 0x12, 0x14) if i == 0 else (
                RGBColor(0xF3, 0xF6, 0xFF) if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name, r.font.size = FONT, Pt(size)
                    r.font.bold = i == 0
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
            c.margin_left = c.margin_right = Emu(70000)
            c.margin_top = c.margin_bottom = Emu(30000)
    return t


def drop(prs, index: int) -> None:
    sld = prs.slides._sldIdLst[index]
    prs.part.drop_rel(sld.rId)
    prs.slides._sldIdLst.remove(sld)
